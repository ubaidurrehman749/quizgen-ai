from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(uploaded_file):
    """
    Extract text from a supported uploaded document.
    Returns the extracted text as a string.
    """

    filename = uploaded_file.name
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file format. "
            "Upload PDF, DOCX, TXT or MD."
        )

    file_bytes = uploaded_file.getvalue()

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))

        if reader.is_encrypted:
            raise ValueError(
                "Password-protected PDFs are not supported."
            )

        pages = []

        for page in reader.pages:
            pages.append(page.extract_text() or "")

        text = "\n".join(pages)

    elif extension == ".docx":
        document = Document(BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        table_text = []

        for table in document.tables:
            for row in table.rows:
                table_text.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        text = "\n".join(paragraphs + table_text)

    else:
        text = file_bytes.decode("utf-8-sig", errors="replace")

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text was found. "
            "The PDF may be scanned or the file may be empty."
        )

    return text
